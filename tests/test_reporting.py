from io import BytesIO
import hashlib
import pandas as pd
import pytest
from openpyxl import load_workbook
from founderos.analytics import load_data, revenue_kpis, support_kpis
from founderos.config import DATA_DIR
from founderos.dashboard import overview, reporting_weeks
from founderos.evidence import dataset_version
from founderos.reporting import REPORTS, template_bytes, prepare_import, ReportingStore
from founderos.periods import compare_metrics, MODES, metric_values, support_records, delivery_records
from founderos.actions import ActionStore

@pytest.fixture
def data(): return load_data(DATA_DIR)

def workbook(report,data,rows=None,mutate=None):
    wb=load_workbook(BytesIO(template_bytes(report,data)))
    for sheet,records in (rows or {}).items():
        headers=[c.value for c in wb[sheet][1]]
        for r,record in enumerate(records,2):
            for c,name in enumerate(headers,1):
                value=record.get(name)
                wb[sheet].cell(r,c,None if pd.isna(value) else value.to_pydatetime() if isinstance(value,pd.Timestamp) else value)
    if mutate: mutate(wb)
    out=BytesIO(); wb.save(out); return out.getvalue()

def ticket(id='REVIEW-TICKET',**kwargs):
    return {'ticket_id':id,'customer_id':'CUS-001','created_at':pd.Timestamp('2026-09-21'),'priority':'high','status':'open','resolved_at':None,'sla_hours':24,'category':'Product issue',**kwargs}

def sales_rows():
    return {'Customers':[{'customer_id':'REVIEW-CUS','company_name':'Reporting demonstration account','segment':'SMB','plan':'Growth','owner':'CS-01','created_at':pd.Timestamp('2026-09-21')}],
            'Subscription Events':[{'event_id':'REVIEW-SUB','customer_id':'REVIEW-CUS','event_date':pd.Timestamp('2026-09-21'),'event_type':'new','mrr_before':0,'mrr_after':900,'reason':'New synthetic subscription'}],
            'Leads':[{'lead_id':'REVIEW-LEAD','created_at':pd.Timestamp('2026-09-21'),'source':'Inbound','status':'qualified'}],
            'Opportunities':[{'opportunity_id':'REVIEW-OPP','lead_id':'REVIEW-LEAD','created_at':pd.Timestamp('2026-09-21'),'stage':'discovery','pipeline_value':10800,'closed_at':None}],
            'Opportunity Stage Events':[{'stage_event_id':'REVIEW-STAGE','opportunity_id':'REVIEW-OPP','event_date':pd.Timestamp('2026-09-21'),'stage':'discovery','pipeline_value':10800}]}

@pytest.mark.parametrize('report',REPORTS)
def test_templates_match_schema_and_have_instructions_validation(data,report):
    content=template_bytes(report,data); wb=load_workbook(BytesIO(content))
    assert wb.sheetnames==['Instructions',*REPORTS[report],'Data Dictionary']
    assert wb['Instructions'].max_row>14 and wb['Data Dictionary'].max_row>2
    assert any(ws.data_validations.count for ws in wb if ws.title in REPORTS[report])
    plan=prepare_import(content,report,data)
    assert not plan.errors and all(f.empty for f in plan.frames.values())

@pytest.mark.parametrize('mutate,field',[
    (lambda w:setattr(w['Tickets'],'title','Wrong tickets'),'Sheet names'),
    (lambda w:w['Tickets'].cell(1,1,'wrong_header'),'Columns'),
    (lambda w:w['Tickets'].cell(2,3,'not a date'),'created_at'),
    (lambda w:w['Tickets'].cell(2,7,'many'),'sla_hours'),
    (lambda w:w['Tickets'].cell(2,2,'CUS-MISSING'),'customer_id'),
    (lambda w:w['Tickets'].cell(2,4,'urgent'),'priority'),
    (lambda w:w['Tickets'].cell(2,1,None),'ticket_id'),
    (lambda w:w['Tickets'].cell(2,3,'2026-09-21T00:00:00Z'),'created_at'),
    (lambda w:w['Tickets'].cell(2,7,'=24+1'),'sla_hours'),
])
def test_invalid_workbook_has_specific_errors(data,mutate,field):
    # openpyxl .cell(..., None) does not clear; clear explicitly for missing ID.
    content=workbook('Customer Support',data,{'Tickets':[ticket()]},mutate)
    if field=='ticket_id':
        content=workbook('Customer Support',data,{'Tickets':[ticket()]},lambda w:setattr(w['Tickets']['A2'],'value',None))
    p=prepare_import(content,'Customer Support',data)
    assert not p.valid and any(e['Field']==field for e in p.errors)
    assert all('Row' in e and 'Worksheet' in e for e in p.errors)

def test_duplicates_and_optional_invalid_dates(data):
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket(),ticket()]}),'Customer Support',data)
    assert any('Duplicate' in e['Message'] for e in p.errors)
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket(resolved_at='bad')]}),'Customer Support',data)
    assert any(e['Field']=='resolved_at' for e in p.errors)

@pytest.mark.parametrize('changed',[
    {'resolved_at':pd.Timestamp('2026-09-20'),'status':'resolved'},
    {'status':'resolved'}, {'sla_hours':0}, {'created_at':pd.Timestamp('2025-01-01')},
])
def test_contradictory_support_rejected(data,changed):
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket(**changed)]}),'Customer Support',data)
    assert p.errors

def test_append_update_persistence_and_recalculation(tmp_path,data):
    store=ReportingStore(tmp_path/'reporting.db',DATA_DIR); before=dataset_version(data)
    content=workbook('Customer Support',data,{'Tickets':[ticket()]})
    plan=prepare_import(content,'Customer Support',store.load(),'support.xlsx')
    assert plan.valid and plan.changes.Operation.tolist()==['Append']
    with pytest.raises(ValueError): store.commit(plan)
    newversion=store.commit(plan,confirmed=True); loaded=ReportingStore(store.path,DATA_DIR).load()
    assert newversion==dataset_version(loaded)!=before
    assert len(loaded['support_tickets'])==len(data['support_tickets'])+1
    assert metric_values(loaded,'2026-09-21')['backlog']==metric_values(data,'2026-09-21')['backlog']+1
    corrected=ticket(status='resolved',resolved_at=pd.Timestamp('2026-09-22'))
    plan=prepare_import(workbook('Customer Support',loaded,{'Tickets':[corrected]}),'Customer Support',loaded)
    assert plan.valid and plan.changes.Operation.tolist()==['Update']
    store.commit(plan,confirmed=True); updated=store.load()
    assert len(updated['support_tickets'])==len(loaded['support_tickets'])
    assert metric_values(updated,'2026-09-21')['backlog']==metric_values(data,'2026-09-21')['backlog']
    assert len(store.history())==2 and store.revision(store.history()[-1]['id'])['support_tickets'].ticket_id.eq('REVIEW-TICKET').any()

def test_sales_consolidates_new_parents_and_changes_kpis(tmp_path,data):
    store=ReportingStore(tmp_path/'reporting.db',DATA_DIR)
    p=prepare_import(workbook('Revenue & Sales',data,sales_rows()),'Revenue & Sales',data,'sales.xlsx')
    assert not p.errors,p.errors
    store.commit(p,confirmed=True); after=store.load(); current=metric_values(after,'2026-09-21'); prior=metric_values(data,'2026-09-21')
    assert current['mrr']==prior['mrr']+900 and current['net_new_mrr']==prior['net_new_mrr']+900
    assert current['active']==prior['active']+1 and current['pipeline']==prior['pipeline']+10800
    assert metric_values(after,'2026-09-14')['mrr']==metric_values(data,'2026-09-14')['mrr']
    assert len(store.history())==1

def test_failed_workbook_is_atomic_and_stale_preview_rejected(tmp_path,data):
    store=ReportingStore(tmp_path/'reporting.db',DATA_DIR)
    bad=sales_rows(); bad['Subscription Events'][0]['customer_id']='MISSING'
    p=prepare_import(workbook('Revenue & Sales',data,bad),'Revenue & Sales',data)
    assert p.errors
    with pytest.raises(ValueError): store.commit(p,confirmed=True)
    assert dataset_version(store.load())==dataset_version(data) and not store.history()
    good=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket()]}),'Customer Support',data)
    store.commit(good,confirmed=True)
    with pytest.raises(ValueError,match='changed'): store.commit(good,confirmed=True)
    assert len(store.history())==1

def test_commit_revalidates_and_sql_failure_rolls_back(tmp_path,data):
    store=ReportingStore(tmp_path/'reporting.db',DATA_DIR)
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket()]}),'Customer Support',data)
    p.frames['support_tickets'].loc[2,'customer_id']='INVALID'
    with pytest.raises(ValueError): store.commit(p,confirmed=True)
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket()]}),'Customer Support',data)
    with store.connect() as c: c.execute("CREATE TRIGGER fail_commit BEFORE INSERT ON reporting_state BEGIN SELECT RAISE(ABORT, 'forced transaction failure'); END")
    with pytest.raises(Exception,match='forced transaction'): store.commit(p,confirmed=True)
    assert not store.history() and dataset_version(store.load())==dataset_version(data)

def test_canonical_restore_preserves_actions_and_source_files(tmp_path,data):
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in DATA_DIR.glob('*.csv')}
    actions=ActionStore(tmp_path/'actions.db'); a=actions.propose('Keep action','Audit','CS-01','2026-10-01',3,dataset_version(data),'2026-09-21'); actions.transition(a['action_id'],'approved')
    original=actions.get(a['action_id']); history=actions.history(a['action_id'])
    store=ReportingStore(tmp_path/'reporting.db',DATA_DIR)
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket()]}),'Customer Support',data); store.commit(p,confirmed=True)
    with pytest.raises(ValueError): store.restore()
    store.restore(confirmed=True)
    for name in data: pd.testing.assert_frame_equal(store.load()[name],data[name],check_dtype=False)
    assert original==actions.get(a['action_id']) and history==actions.history(a['action_id'])
    assert len(store.history())==2
    assert hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in DATA_DIR.glob('*.csv')}

def test_mrr_chain_and_missing_stage_history_rejected(data):
    rows=sales_rows(); rows['Subscription Events'][0]['mrr_before']=100
    p=prepare_import(workbook('Revenue & Sales',data,rows),'Revenue & Sales',data)
    assert any('MRR' in e['Message'] for e in p.errors)
    rows=sales_rows(); rows['Opportunity Stage Events']=[]
    p=prepare_import(workbook('Revenue & Sales',data,rows),'Revenue & Sales',data)
    assert any('stage events' in e['Message'] for e in p.errors)

@pytest.mark.parametrize('mode',MODES)
def test_period_comparisons_use_each_cutoff(data,mode):
    rows,dates=compare_metrics(data,pd.Timestamp('2026-09-21'),mode,'2026-07-06')
    mrr=next(r for r in rows if r['key']=='mrr')
    expected=sum(metric_values(data,w)['mrr'] for w in dates)/len(dates)
    assert mrr['Comparison']==expected and mrr['Change']==pytest.approx(mrr['Current']-expected)
    assert dates==[pd.Timestamp('2026-08-24')] if mode==MODES[1] else True

def test_insufficient_history_and_zero_denominators(data):
    rows,dates=compare_metrics(data,'2026-07-06',MODES[2]); assert not dates and all(r['Comparison'] is None for r in rows)
    assert metric_values(data,'2026-09-21')['win'] is None
    zero={n:f.head(0).copy() for n,f in data.items()}
    for k in ['gross_revenue_churn','sla','resolution_time','on_time','capacity','win']: assert metric_values(zero,'2026-09-21')[k] is None

def test_sunday_mrr_and_future_resolutions_cutoff(data):
    event=pd.DataFrame([{'event_id':'EDGE','customer_id':'EDGE','event_date':pd.Timestamp('2026-09-27 23:59:59'),'event_type':'new','mrr_before':0,'mrr_after':123,'reason':'Boundary'}])
    assert revenue_kpis(event,'2026-09-21')['mrr']==123
    event.loc[0,'event_date']=pd.Timestamp('2026-09-28'); assert revenue_kpis(event,'2026-09-21')['mrr']==0
    d={n:f.copy() for n,f in data.items()}; d['support_tickets']=pd.DataFrame([ticket(resolved_at=pd.Timestamp('2026-09-28'),status='resolved')]); d['support_tickets']['resolved_at']=pd.to_datetime(d['support_tickets'].resolved_at)
    k=support_kpis(d['support_tickets'],'2026-09-21'); assert k['tickets_resolved']==0 and k['backlog']==1 and k['open_sla_breaches']==1
    row=support_records(d,'2026-09-21').iloc[0]; assert row.age_hours==168 and row.hours_past_deadline==144 and row.sla_deadline==pd.Timestamp('2026-09-22') and pd.isna(row.resolved_at)
    assert set(d['support_tickets'].category)=={'Product issue'} # no complaint classification inferred

def test_future_delivery_completion_hidden(data):
    d={n:f.copy() for n,f in data.items()}; d['implementation_milestones'].loc[0,'completed_at']=pd.Timestamp('2026-10-01')
    _,m=delivery_records(d,'2026-09-21'); row=m[m.milestone_id.eq(d['implementation_milestones'].loc[0,'milestone_id'])].iloc[0]
    assert pd.isna(row.completed_at) and row.status=='open' and row.overdue

def test_support_import_changes_alerts(tmp_path,data):
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket('DETERIORATION-'+str(i)) for i in range(25)]}),'Customer Support',data)
    s=ReportingStore(tmp_path/'r.db',DATA_DIR); s.commit(p,confirmed=True)
    old=overview(data,pd.Timestamp('2026-09-21')); new=overview(s.load(),pd.Timestamp('2026-09-21'))
    assert new['snapshot']['support']['ticket_volume']==old['snapshot']['support']['ticket_volume']+25
    alert=next(a for a in new['alerts'] if a['alert_id']=='support_backlog'); assert alert['evidence']['current_backlog']==109

def test_weekly_capacity_composite_key_update_and_validation(tmp_path,data):
    row=data['team_capacity'].iloc[0].to_dict(); row['allocated_hours']=39
    p=prepare_import(workbook('Team Capacity',data,{'Weekly Capacity':[row]}),'Team Capacity',data)
    assert not p.errors and p.changes.Operation.tolist()==['Update']
    s=ReportingStore(tmp_path/'r.db',DATA_DIR); s.commit(p,confirmed=True)
    loaded=s.load(); assert len(loaded['team_capacity'])==len(data['team_capacity'])
    assert metric_values(loaded,row['week_start'])['capacity']!=metric_values(data,row['week_start'])['capacity']
    row['function']='Sales' if row['function']!='Sales' else 'Management'
    p=prepare_import(workbook('Team Capacity',data,{'Weekly Capacity':[row]}),'Team Capacity',data)
    assert any('match the team' in e['Message'] for e in p.errors)

def test_new_activity_week_without_capacity_is_available(tmp_path,data):
    p=prepare_import(workbook('Customer Support',data,{'Tickets':[ticket(created_at=pd.Timestamp('2026-09-28'))]}),'Customer Support',data)
    s=ReportingStore(tmp_path/'r.db',DATA_DIR); s.commit(p,confirmed=True)
    loaded=s.load(); assert reporting_weeks(loaded)[-1]==pd.Timestamp('2026-09-28')
    assert metric_values(loaded,'2026-09-28')['capacity'] is None
    assert metric_values(loaded,'2026-09-21')['backlog']==metric_values(data,'2026-09-21')['backlog']
    assert 'complaints' not in metric_values(loaded,'2026-09-28')

def test_exact_sla_boundary_and_on_time_vs_late_completion(data):
    d={n:f.copy() for n,f in data.items()}
    d['support_tickets']=pd.DataFrame([ticket(created_at=pd.Timestamp('2026-09-27'),sla_hours=24)])
    d['support_tickets']['resolved_at']=pd.to_datetime(d['support_tickets'].resolved_at)
    assert support_records(d,'2026-09-21').iloc[0].hours_past_deadline==0
    assert metric_values(d,'2026-09-21')['breaches']==0
    m=d['implementation_milestones'].head(2).copy(); m['due_date']=pd.Timestamp('2026-09-23'); m['completed_at']=[pd.Timestamp('2026-09-23'),pd.Timestamp('2026-09-24')]; d['implementation_milestones']=m
    values=metric_values(d,'2026-09-21'); assert values['on_time']==.5 and values['completed']==2

def test_parent_correction_cannot_invalidate_retained_children(data):
    row=data['customers'].iloc[0].to_dict(); row['created_at']=pd.Timestamp('2026-10-01')
    p=prepare_import(workbook('Revenue & Sales',data,{'Customers':[row]}),'Revenue & Sales',data)
    assert any('Cannot precede' in e['Message'] for e in p.errors)
