"""Transactional reporting workspace and schema-aligned Excel reports."""
from __future__ import annotations
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import hashlib
import json
import sqlite3
import zipfile
from contextlib import contextmanager
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from .analytics import DATE_COLUMNS, load_data
from .dashboard import REQUIRED_COLUMNS, ID_COLUMN, FK_RULES
from .evidence import dataset_version

REPORTS = {
    'Revenue & Sales': {'Customers':'customers', 'Subscription Events':'subscription_events', 'Financial Entries':'financial_entries', 'Leads':'leads', 'Opportunities':'opportunities', 'Opportunity Stage Events':'opportunity_stage_events'},
    'Customer Support': {'Tickets':'support_tickets'},
    'Delivery': {'Projects':'implementation_projects', 'Milestones':'implementation_milestones'},
    'Team Capacity': {'Team Members':'team', 'Weekly Capacity':'team_capacity'},
}
REPORT_LABELS={'Revenue & Sales':'Revenue & Sales','Customer Support':'Customer Support','Delivery':'Customer Onboarding & Projects','Team Capacity':'Team Workload'}
ENUMS = {
    ('customers','segment'):['SMB','Mid-market'], ('customers','plan'):['Starter','Growth','Scale'],
    ('subscription_events','event_type'):['opening','new','expansion','contraction','churn'],
    ('leads','status'):['qualified','disqualified'], ('leads','source'):['Inbound','Outbound','Partner'],
    ('opportunities','stage'):['discovery','proposal','won','lost'],
    ('opportunity_stage_events','stage'):['discovery','proposal','won','lost'],
    ('support_tickets','priority'):['low','medium','high'], ('support_tickets','status'):['open','resolved'],
    ('support_tickets','category'):['How-to','Product issue'],
    ('implementation_projects','status'):['active','complete'], ('implementation_milestones','status'):['open','complete'],
    ('implementation_milestones','milestone'):['Kickoff','Configuration','Training','Go-live'],
    ('team','function'):['Management','Sales','Customer Success','Product','Engineering'],
    ('team_capacity','function'):['Management','Sales','Customer Success','Product','Engineering'],
    ('financial_entries','entry_type'):['recognized_revenue','direct_cost','cash_receipt','cash_payment','opening_cash','closing_cash'],
}
NUMBERS = {'mrr_before','mrr_after','pipeline_value','sla_hours','allocated_hours','capacity_hours','weekly_capacity_hours','amount'}
OPTIONAL = {'closed_at','resolved_at','completed_at'}
DESCRIPTIONS = {
    'customer_id':'Stable customer ID; must exist in Customers for dependent records.',
    'company_name':'Customer business name.', 'segment':'Customer segment.', 'plan':'Subscription plan.',
    'owner':'Accountable employee ID from Team Members.', 'created_at':'Record creation timestamp.',
    'event_id':'Unique subscription event ID.', 'event_date':'Effective event date/time; drives historical cutoffs.',
    'event_type':'Subscription movement type.', 'mrr_before':'Monthly recurring EUR before the event.',
    'mrr_after':'Monthly recurring EUR after the event.', 'reason':'Business reason for the movement.',
    'lead_id':'Stable lead ID; must exist in Leads for opportunities.', 'source':'Lead acquisition channel.',
    'status':'Current operational state; dates drive historical state where available.',
    'opportunity_id':'Stable opportunity ID.', 'stage':'Sales lifecycle stage.',
    'pipeline_value':'Unweighted contract opportunity value in EUR; not MRR.', 'closed_at':'Closing timestamp; blank for open opportunities.',
    'stage_event_id':'Unique opportunity lifecycle event ID; required for historical pipeline.',
    'ticket_id':'Stable support ticket ID.', 'priority':'Support urgency.', 'resolved_at':'Resolution timestamp; blank if unresolved.',
    'sla_hours':'Resolution target in elapsed calendar hours from creation.', 'category':'Support theme; neither category establishes a complaint.',
    'project_id':'Stable implementation project ID.', 'start_date':'Implementation start date.',
    'target_go_live':'Committed go-live date.', 'milestone_id':'Stable milestone ID.', 'milestone':'Implementation commitment type.',
    'due_date':'Milestone deadline; date-only values mean midnight at the start of that date.',
    'completed_at':'Actual completion timestamp; blank while incomplete.',
    'employee_id':'Stable team member ID.', 'function':'Team function; must match Team Members.', 'role':'Team member role.',
    'weekly_capacity_hours':'Nominal available weekly hours, greater than zero.',
    'week_start':'Monday reporting date; together with employee_id forms the stable key.',
    'allocated_hours':'Allocated hours for this reporting week; overload above capacity is allowed.',
    'capacity_hours':'Available hours for this week, greater than zero.',
    'entry_id':'Stable unique financial journal ID; reuse it to correct a posted entry.',
    'entry_date':'Effective financial date. Recognized revenue is posted at month end; bank closing balances use the last calendar day.',
    'entry_type':'Recognized subscription service, direct service cost, cash movement or bank statement balance. Balances are not cash movements.',
    'amount':'EUR value. Revenue, service costs and movements are nonnegative; opening/closing bank balances may be signed.',
    'description':'Business description. Closing balances come from the bank statement, not a manually calculated KPI.',
}

def nullable(name,col):
    return col in OPTIONAL or (name=='financial_entries' and col=='customer_id')

def keys(name):
    return ['week_start','employee_id'] if name=='team_capacity' else [ID_COLUMN[name]]

def _json_data(data):
    return json.dumps({name:json.loads(frame.to_json(orient='records',date_format='iso')) for name,frame in data.items()},allow_nan=False)

def _decode(encoded):
    raw=json.loads(encoded)
    data={name:pd.DataFrame(raw.get(name,[]),columns=cols) for name,cols in REQUIRED_COLUMNS.items()}
    for name,cols in DATE_COLUMNS.items():
        for col in cols: data[name][col]=pd.to_datetime(data[name][col],format='mixed')
    for frame in data.values():
        for col in NUMBERS.intersection(frame.columns): frame[col]=pd.to_numeric(frame[col])
    return data

def template_bytes(report, data):
    """Blank import tables; real example rows live in Instructions, not import sheets."""
    wb=Workbook(); instructions=wb.active; instructions.title='Instructions'
    lines=[['RelayOps '+REPORT_LABELS[report]+' reporting'], ['Complete the data worksheets; do not rename sheets or headers.'],
           ['Submit new IDs to append; submit existing IDs to correct the entire row. No deletion by omission.'],
           ['For corrections, include every required field. Blank optional dates clear the previous value.'],
           ['Keep IDs stable and unique. Weekly Capacity uses week_start plus employee_id.'],
           ['Use YYYY-MM-DD or YYYY-MM-DD HH:MM:SS; all timestamps use one consistent business timezone, without offsets.'],
           ['Weeks run Monday to the next Monday exclusive. Cutoff is next Monday at 00:00.'],
           ['Upload, inspect row errors and preview, then explicitly confirm or cancel.'],
           ['Required means every column is required; closed_at/resolved_at/completed_at may be blank.'],
           ['Existing customer/team IDs must be used; include new parents in the same workbook when supported.'],
           ['Subscription and opportunity events are the historical sources; keep their full lifecycle coherent.'],
           ['Lead status has no qualification timestamp; conversion is a created-lead cohort proxy.'],
           ['Tickets are not automatically complaints. Examples below are reference only; they are not imported.'], []]
    if report=='Revenue & Sales':
        lines.insert(13,['Financial Entries: monthly subscription service accrual, direct costs and cash movements are distinct. Use bank statement balances; FounderOS checks reconciliation.'])
        lines.insert(14,['Revenue is prorated by effective subscription prices over calendar-month service time. Opening subscription balances carry from the start of their opening month.'])
        lines.insert(15,['Finance customer_id is required for recognized subscription revenue/receipts; leave it blank for other entries. One bank opening and one month-end closing balance per closed month.'])
    for line in lines: instructions.append(line)
    dictionary=wb.create_sheet('Data Dictionary'); dictionary.append(['Worksheet','Field','Required','Description','Format / accepted values','Stable key'])
    for sheet,name in REPORTS[report].items():
        ws=wb.create_sheet(sheet); ws.append(REQUIRED_COLUMNS[name]); ws.freeze_panes='B2'; ws.auto_filter.ref=f'A1:{get_column_letter(ws.max_column)}1001'
        instructions.append([sheet+' example (existing record; submitting it updates that ID)'])
        instructions.append(REQUIRED_COLUMNS[name])
        row=data[name].head(1)
        if len(row): instructions.append([None if pd.isna(v) else v.to_pydatetime() if isinstance(v,pd.Timestamp) else v.item() if hasattr(v,'item') else v for v in row.iloc[0]])
        instructions.append([])
        for i,col in enumerate(REQUIRED_COLUMNS[name],1):
            accepted=ENUMS.get((name,col)); date=col in DATE_COLUMNS.get(name,[])
            fmt=', '.join(accepted) if accepted else 'YYYY-MM-DD HH:MM:SS' if date else 'EUR; bank opening/closing balances may be signed' if col=='amount' else 'Nonnegative number (EUR / hours as described)' if col in NUMBERS else 'Text; no formulas'
            dictionary.append([sheet,col,'Column required; blank allowed' if nullable(name,col) else 'Required value',DESCRIPTIONS[col],fmt,'Yes' if col in keys(name) else ''])
            ws.cell(1,i).comment=Comment(DESCRIPTIONS[col]+' '+fmt,'RelayOps')
            ws.column_dimensions[get_column_letter(i)].width=max(18,min(29,len(col)+5))
            if accepted:
                rule=DataValidation(type='list',formula1='"'+','.join(accepted)+'"'); rule.errorTitle='Choose an accepted value'; rule.error='Use the dropdown values.'; rule.showErrorMessage=True; rule.errorStyle='stop'; ws.add_data_validation(rule); rule.add(f'{get_column_letter(i)}2:{get_column_letter(i)}1001')
            for r in range(2,102):
                ws.cell(r,i).number_format='yyyy-mm-dd hh:mm' if date else '#,##0.00' if col in NUMBERS else '@'
                ws.cell(r,i).fill=PatternFill('solid',fgColor='FFF8DF')
    wb._sheets=[instructions]+[wb[sheet] for sheet in REPORTS[report]]+[dictionary]
    for ws in wb:
        ws.sheet_view.showGridLines=False
        for cell in ws[1]: cell.font=Font(name='Arial',bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='193252'); cell.alignment=Alignment(vertical='center')
        ws.row_dimensions[1].height=26
        if ws.title=='Data Dictionary':
            ws.freeze_panes='C2'
            for col,width in zip('ABCDEF',[28,28,30,70,66,16]): ws.column_dimensions[col].width=width
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top'); cell.font=Font(name='Arial',size=10)
            for r in range(2,ws.max_row+1): ws.row_dimensions[r].height=32
        elif ws.title=='Instructions':
            ws.column_dimensions['A'].width=110
            for col in range(2,12): ws.column_dimensions[get_column_letter(col)].width=23
    out=BytesIO(); wb.save(out); return out.getvalue()

@dataclass
class ImportPlan:
    report: str
    base_version: str
    frames: dict
    changes: pd.DataFrame
    errors: list
    filename: str = ''
    @property
    def valid(self): return not self.errors

def parse_workbook(content,report):
    frames={}; errors=[]
    try:
        if len(content)>10_000_000: raise ValueError('Workbook exceeds 10 MB.')
        with zipfile.ZipFile(BytesIO(content)) as package:
            if sum(x.file_size for x in package.infolist())>50_000_000: raise ValueError('Expanded workbook exceeds 50 MB.')
        wb=load_workbook(BytesIO(content),data_only=False,read_only=True)
        expected={'Instructions','Data Dictionary',*REPORTS[report]}
        if set(wb.sheetnames)!=expected:
            errors.append({'Worksheet':'Workbook','Row':'—','Field':'Sheet names','Message':f'Expected exactly: {", ".join(expected)}. Found: {", ".join(wb.sheetnames)}.'})
            return frames,errors
        for sheet,name in REPORTS[report].items():
            ws=wb[sheet]
            if ws.max_row>10001: raise ValueError(f'{sheet} exceeds 10,000 records.')
            rows=ws.iter_rows(values_only=True); headers=list(next(rows,()))
            while headers and headers[-1] is None: headers.pop()
            if len(headers)!=len(set(headers)) or set(headers)!=set(REQUIRED_COLUMNS[name]):
                errors.append({'Worksheet':sheet,'Row':1,'Field':'Columns','Message':'Use the template headers exactly. Required: '+', '.join(REQUIRED_COLUMNS[name])}); continue
            values=[]; indexes=[]
            for r,row in enumerate(rows,2):
                if not any(v is not None for v in row): continue
                values.append(dict(zip(headers,row))); indexes.append(r)
            frames[name]=pd.DataFrame(values,columns=REQUIRED_COLUMNS[name],index=indexes)
        wb.close()
    except Exception as exc:
        errors.append({'Worksheet':'Workbook','Row':'—','Field':'File','Message':f'Cannot read Excel report: {exc}'})
    return frames,errors

def prepare_import(content,report,current,filename=''):
    frames,errors=parse_workbook(content,report); return prepare_frames(frames,report,current,errors,filename)

def prepare_frames(frames,report,current,errors=None,filename=''):
    errors=list(errors or []); typed={}; changes=[]; combined={n:f.copy() for n,f in current.items()}
    inverse={n:s for s,n in REPORTS[report].items()}
    def error(name,r,col,message): errors.append({'Worksheet':inverse.get(name,name),'Row':r,'Field':col,'Message':message})
    for name,frame in frames.items():
        f=frame.copy(); typed[name]=f
        for col in f:
            for r,v in f[col].items():
                if isinstance(v,str) and v.startswith(('=','+','@')): error(name,r,col,'Formulas and executable expressions are not accepted.'); continue
                if pd.isna(v) or (isinstance(v,str) and not v.strip()):
                    f.loc[r,col]=None
                    if not nullable(name,col): error(name,r,col,'A value is required.')
        for col in DATE_COLUMNS.get(name,[]):
            original=f[col].copy()
            parsed=original.map(lambda v: pd.to_datetime(v,errors='coerce') if pd.notna(v) and not isinstance(v,(int,float,bool)) else pd.NaT)
            for r,v in original.items():
                if pd.notna(v) and (isinstance(v,(int,float)) or pd.isna(parsed.loc[r])): error(name,r,col,'Use a valid Excel date or YYYY-MM-DD timestamp, not a numeric serial.')
                if pd.notna(parsed.loc[r]) and parsed.loc[r].tzinfo is not None:
                    error(name,r,col,'Use a consistent business time without a timezone offset.'); parsed.loc[r]=pd.NaT
            f[col]=pd.to_datetime(parsed)
        for col in NUMBERS.intersection(f.columns):
            parsed=pd.to_numeric(f[col],errors='coerce')
            for r,v in parsed.items():
                if isinstance(f.loc[r,col],bool) or pd.isna(v) or not float('-inf')<v<float('inf'): error(name,r,col,'Use a finite number.')
                elif (v<0 and not (name=='financial_entries' and col=='amount' and f.loc[r,'entry_type'] in ['opening_cash','closing_cash'])) or (col in {'sla_hours','capacity_hours','weekly_capacity_hours'} and v==0): error(name,r,col,'Must be greater than zero.' if col in {'sla_hours','capacity_hours','weekly_capacity_hours'} else 'Cannot be negative.')
            f[col]=parsed
        for col in set(f.columns)-NUMBERS-set(DATE_COLUMNS.get(name,[])):
            for r,v in f[col].items():
                if pd.notna(v) and not isinstance(v,str): error(name,r,col,'Use text, preserving stable IDs.')
                elif isinstance(v,str): f.loc[r,col]=v.strip()
        for (table,col),values in ENUMS.items():
            if table==name:
                for r,v in f[col].items():
                    if v not in values: error(name,r,col,'Accepted values: '+', '.join(values))
        for r in f.index[f.duplicated(keys(name),keep=False)]: error(name,r,', '.join(keys(name)),'Duplicate identifier in this workbook; submit each key once.')
        if name=='team_capacity':
            for r,v in f.week_start.items():
                if pd.notna(v) and (v.dayofweek!=0 or v!=v.normalize()): error(name,r,'week_start','Use a Monday date at midnight.')
        if errors: continue
        old=current[name].set_index(keys(name),drop=False); new=f.set_index(keys(name),drop=False)
        for key,row in new.iterrows():
            action='Append' if key not in old.index else 'Unchanged' if row.equals(old.loc[key]) else 'Update'
            # Compare values independent of dataframe dtype and null representation.
            if key in old.index:
                action='Unchanged' if all((pd.isna(row[c]) and pd.isna(old.loc[key,c])) or row[c]==old.loc[key,c] for c in f.columns) else 'Update'
            fields=[] if action!='Update' else [c for c in f.columns if not ((pd.isna(row[c]) and pd.isna(old.loc[key,c])) or row[c]==old.loc[key,c])]
            before={c:None if pd.isna(old.loc[key,c]) else str(old.loc[key,c]) for c in fields}
            after={c:None if pd.isna(row[c]) else str(row[c]) for c in fields}
            changes.append({'Worksheet':inverse[name],'Row':_source_row(f,keys(name),row), 'Record ID':str(key),'Operation':action,'Changed fields':', '.join(fields),'Previous values':json.dumps(before) if fields else '', 'Submitted values':json.dumps(after) if fields else ''})
        remove=set(new.index); retained=old.loc[[key not in remove for key in old.index]]
        combined[name]=pd.concat([retained.reset_index(drop=True),f.reset_index(drop=True)],ignore_index=True)
    if not errors:
        # Cross-table rules run on the consolidated candidate, including new parents.
        for name,f in typed.items():
            for r,row in f.iterrows():
                if name in FK_RULES:
                    col,parent,pkey=FK_RULES[name]
                    if row[col] not in set(combined[parent][pkey]): error(name,r,col,f'Unknown {parent} ID: {row[col]}.')
                if name=='financial_entries':
                    if row.entry_type in ['recognized_revenue','cash_receipt'] and pd.isna(row.customer_id): error(name,r,'customer_id','Customer reference is required for subscription revenue and receipts.')
                    if pd.notna(row.customer_id) and row.customer_id not in set(combined['customers'].customer_id): error(name,r,'customer_id','Unknown customer ID in financial report.')
                    if pd.notna(row.customer_id) and row.customer_id in set(combined['customers'].customer_id):
                        created=combined['customers'].set_index('customer_id').loc[row.customer_id,'created_at']
                        if row.entry_date.normalize()<created.normalize(): error(name,r,'entry_date','Financial customer entries cannot precede the customer creation date.')
                    if row.entry_type=='recognized_revenue' and row.entry_date.normalize()!=row.entry_date.to_period('M').end_time.normalize(): error(name,r,'entry_date','Monthly recognized revenue must be posted on the last day of the month.')
                if 'owner' in f and row.owner not in set(combined['team'].employee_id): error(name,r,'owner','Unknown team member ID.')
                pairs={'support_tickets':('created_at','resolved_at'),'opportunities':('created_at','closed_at'),'implementation_projects':('start_date','target_go_live')}
                if name in pairs:
                    start,end=pairs[name]
                    if pd.notna(row[end]) and row[end]<row[start]: error(name,r,end,f'Cannot precede {start}.')
                if name in {'support_tickets','implementation_milestones'}:
                    date='resolved_at' if name=='support_tickets' else 'completed_at'; closed='resolved' if name=='support_tickets' else 'complete'
                    if (row.status==closed)!=pd.notna(row[date]): error(name,r,date,f'{closed} status requires a completion date; open status requires blank.')
                if name=='opportunities' and (row.stage in ['won','lost'])!=pd.notna(row.closed_at): error(name,r,'closed_at','Won/lost needs a closing date; discovery/proposal needs blank.')
                if name=='team_capacity' and row.employee_id in set(combined['team'].employee_id):
                    member=combined['team'].set_index('employee_id').loc[row.employee_id]
                    if row.function!=member.function: error(name,r,'function','Must match the team member function.')
                if name=='implementation_milestones' and row.project_id in set(combined['implementation_projects'].project_id):
                    start=combined['implementation_projects'].set_index('project_id').loc[row.project_id,'start_date']
                    if row.due_date<start or (pd.notna(row.completed_at) and row.completed_at<start): error(name,r,'due_date / completed_at','Milestone dates cannot precede project start.')
        # A correction to a parent cannot invalidate retained child rows.
        for name,(col,parent,pkey) in FK_RULES.items():
            if not set(combined[name][col]).issubset(set(combined[parent][pkey])): error(name,'—',col,'Consolidated records contain an invalid reference.')
        for name in ['customers','implementation_projects']:
            if not set(combined[name].owner).issubset(set(combined['team'].employee_id)): error(name,'—','owner','Unknown accountable team member in consolidated records.')
        for name,date,parent,pkey,parentdate in [('subscription_events','event_date','customers','customer_id','created_at'),('support_tickets','created_at','customers','customer_id','created_at'),('opportunities','created_at','leads','lead_id','created_at'),('opportunity_stage_events','event_date','opportunities','opportunity_id','created_at'),('implementation_projects','start_date','customers','customer_id','created_at')]:
            joined=combined[name].merge(combined[parent][[pkey,parentdate]].rename(columns={parentdate:'_parent_date'}),on=pkey)
            bad=joined[joined[date]<joined._parent_date]
            prior=current[name].merge(current[parent][[pkey,parentdate]].rename(columns={parentdate:'_parent_date'}),on=pkey)
            legacy=set(prior.loc[prior[date]<prior._parent_date,keys(name)[0]])
            submitted=set(typed[name][keys(name)[0]]) if name in typed else set()
            for _,row in bad.iterrows():
                if row[keys(name)[0]] not in legacy or row[keys(name)[0]] in submitted: error(name,_source_row(typed.get(name),keys(name),row),date,f'Cannot precede referenced {parent} creation.')
        events=combined['subscription_events']
        for _,group in events.sort_values(['customer_id','event_date','event_id']).groupby('customer_id'):
            previous=float(group.iloc[0].mrr_before) if group.iloc[0].event_type=='opening' else 0.0
            for position,(_,row) in enumerate(group.iterrows()):
                if abs(row.mrr_before-previous)>.005: error('subscription_events',_source_row(typed.get('subscription_events'),['event_id'],row),'mrr_before',f'Contradictory MRR history for {row.customer_id}; expected {previous:g}.')
                direction=row.mrr_after-row.mrr_before
                valid={'opening':position==0 and row.mrr_after>0 and row.mrr_before in [0,row.mrr_after],'new':row.mrr_before==0 and row.mrr_after>0,'expansion':direction>0 and row.mrr_before>0,'contraction':direction<0 and row.mrr_after>0,'churn':row.mrr_before>0 and row.mrr_after==0}
                if not valid[row.event_type]: error('subscription_events',_source_row(typed.get('subscription_events'),['event_id'],row),'event_type','Movement type contradicts before/after MRR.')
                previous=row.mrr_after
        cap=combined['team_capacity'].merge(combined['team'][['employee_id','function']],on='employee_id',suffixes=('','_team'))
        if cap.function.ne(cap.function_team).any(): error('team_capacity','—','function','Team correction conflicts with retained weekly capacity functions.')
        # Validate edited opportunities and stage histories together. Legacy current-state
        # fields are not used to reconstruct pipeline; imported pairs must be coherent.
        changed_opps=set(typed.get('opportunities',pd.DataFrame(columns=['opportunity_id'])).opportunity_id)|set(typed.get('opportunity_stage_events',pd.DataFrame(columns=['opportunity_id'])).opportunity_id)
        for oid in changed_opps:
            opportunity=combined['opportunities'].set_index('opportunity_id').loc[oid] if oid in set(combined['opportunities'].opportunity_id) else None
            if opportunity is None: continue
            lifecycle=combined['opportunity_stage_events'][combined['opportunity_stage_events'].opportunity_id.eq(oid)].sort_values(['event_date','stage_event_id'])
            if lifecycle.empty: error('opportunities',_source_row(typed.get('opportunities'),['opportunity_id'],{'opportunity_id':oid}),'opportunity_id','Include dated stage events; otherwise pipeline cannot be calculated.')
            else:
                if lifecycle.event_date.duplicated().any(): error('opportunity_stage_events','—','event_date',f'Ambiguous same-time stage events for {oid}.')
                latest=lifecycle.iloc[-1]
                if latest.stage!=opportunity.stage or abs(latest.pipeline_value-opportunity.pipeline_value)>.005: error('opportunities',_source_row(typed.get('opportunities'),['opportunity_id'],{'opportunity_id':oid}),'stage / pipeline_value','Current opportunity must match its latest stage event.')
                if opportunity.stage in ['won','lost'] and latest.event_date!=opportunity.closed_at: error('opportunities','—','closed_at','Closing timestamp must match the won/lost stage event.')
        # Parent corrections must not move an existing milestone before project start.
        all_m=combined['implementation_milestones'].merge(combined['implementation_projects'][['project_id','start_date']],on='project_id')
        for _,row in all_m[(all_m.due_date<all_m.start_date)|(all_m.completed_at.notna()&(all_m.completed_at<all_m.start_date))].iterrows(): error('implementation_milestones',_source_row(typed.get('implementation_milestones'),['milestone_id'],row),'due_date / completed_at','Milestone cannot precede project start after consolidation.')
        if 'financial_entries' in typed and len(typed['financial_entries']):
            from .business import finance_errors
            for record,message in finance_errors(combined):
                selected=typed['financial_entries'][typed['financial_entries'].entry_id.eq(record)]
                error('financial_entries',selected.index[0] if len(selected) else 'Consolidated month / customer '+str(record),'amount / reconciliation',message)
    return ImportPlan(report,dataset_version(current),typed,pd.DataFrame(changes),errors,Path(filename).name)

def _source_row(frame,columns,row):
    if frame is None: return 'Existing record '+str(tuple(row[c] for c in columns))
    mask=pd.Series(True,index=frame.index)
    for col in columns: mask &= frame[col].eq(row[col])
    return frame.index[mask][0] if mask.any() else 'Existing record '+str(tuple(row[c] for c in columns))

def apply_frames(current,frames):
    updated={n:f.copy() for n,f in current.items()}
    for name,f in frames.items():
        old=updated[name].set_index(keys(name),drop=False); incoming=f.set_index(keys(name),drop=False)
        updated[name]=pd.concat([old.loc[[k not in incoming.index for k in old.index]].reset_index(drop=True),f.reset_index(drop=True)],ignore_index=True)
    return updated

class ReportingStore:
    """Whole-dataset immutable revisions committed atomically in a separate SQLite DB."""
    def __init__(self,path,canonical_dir):
        self.path=Path(path); self.canonical_dir=Path(canonical_dir); self.path.parent.mkdir(parents=True,exist_ok=True)
        with self.connect() as c:
            c.executescript('CREATE TABLE IF NOT EXISTS revisions (id INTEGER PRIMARY KEY AUTOINCREMENT, version TEXT, data_json TEXT, report TEXT, filename TEXT, summary_json TEXT, imported_at TEXT); CREATE TABLE IF NOT EXISTS reporting_state (singleton INTEGER PRIMARY KEY CHECK(singleton=1), revision_id INTEGER);')
    @contextmanager
    def connect(self):
        connection=sqlite3.connect(self.path,timeout=15)
        try:
            with connection: yield connection
        finally: connection.close()
    def load(self):
        with self.connect() as c:
            row=c.execute('SELECT data_json FROM revisions JOIN reporting_state ON revisions.id=reporting_state.revision_id WHERE singleton=1').fetchone()
        return _decode(row[0]) if row else load_data(self.canonical_dir)
    def commit(self,plan,confirmed=False):
        if confirmed is not True: raise ValueError('Confirm the preview before importing.')
        with self.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            row=c.execute('SELECT data_json FROM revisions JOIN reporting_state ON revisions.id=reporting_state.revision_id WHERE singleton=1').fetchone()
            current=_decode(row[0]) if row else load_data(self.canonical_dir)
            if dataset_version(current)!=plan.base_version: raise ValueError('Reporting data changed; validate the workbook again.')
            checked=prepare_frames(plan.frames,plan.report,current,filename=plan.filename)
            if plan.errors or checked.errors: raise ValueError('Workbook validation failed; no records were changed.')
            if checked.changes.empty or checked.changes.Operation.eq('Unchanged').all(): raise ValueError('No new or changed records to import.')
            updated=apply_frames(current,checked.frames)
            return self._save(c,updated,plan.report,plan.filename,checked.changes.to_dict('records'))
    def _save(self,c,data,report,filename,summary):
        encoded=_json_data(data); normalized=_decode(encoded); version=dataset_version(normalized)
        cursor=c.execute('INSERT INTO revisions(version,data_json,report,filename,summary_json,imported_at) VALUES(?,?,?,?,?,?)',(version,encoded,report,filename,json.dumps(summary,default=str),datetime.now(timezone.utc).isoformat()))
        c.execute('INSERT INTO reporting_state VALUES(1,?) ON CONFLICT(singleton) DO UPDATE SET revision_id=excluded.revision_id',(cursor.lastrowid,))
        return version
    def restore(self,confirmed=False):
        if confirmed is not True: raise ValueError('Explicit confirmation is required to restore demonstration data.')
        with self.connect() as c:
            c.execute('BEGIN IMMEDIATE'); return self._save(c,load_data(self.canonical_dir),'Restore demonstration data','',[])
    def history(self):
        with self.connect() as c:
            c.row_factory=sqlite3.Row
            return [dict(r) for r in c.execute('SELECT id,version,report,filename,summary_json,imported_at FROM revisions ORDER BY id DESC')]
    def revision(self,revision_id):
        with self.connect() as c: row=c.execute('SELECT data_json FROM revisions WHERE id=?',(revision_id,)).fetchone()
        if not row: raise ValueError('Unknown reporting revision.')
        return _decode(row[0])
